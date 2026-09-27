import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("AEGIS | MAJOR PROJECT REVIEW 2 IMPLEMENTATION PLAN\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5) # Indigo

    run_title = title_p.add_run("Autonomous Engineering Graph & Intelligence System\nPrototype Implementation & Engineering Specification")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate 900

    # Motto Banner
    motto_table = doc.add_table(rows=1, cols=1)
    motto_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = motto_table.cell(0, 0)
    set_cell_background(cell, "EEF2FF") # Indigo 50
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    mp = cell.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m_run = mp.add_run("PROJECT MOTTO:\n\"See every risk before it ships. Heal every failure after it lands.\"")
    m_run.font.bold = True
    m_run.font.size = Pt(12)
    m_run.font.color.rgb = RGBColor(0x37, 0x30, 0xA3) # Indigo 800

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Metadata Block
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Deliverable Target:", "Review 2 Interactive Mission Control Prototype"), ("Architecture:", "Frontend + Minimal Backend + Real Algorithms + Real AI")],
        [("Substrate:", "Neo4j Knowledge Graph / Graph In-Memory + LLM RCA Agent"), ("Safety Invariant:", "Strict Decoupling of Reasoning Layer & Guarded Actuation")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (k, v) in enumerate(row):
            c = meta_table.cell(r_idx, c_idx)
            set_cell_background(c, "F8FAFC")
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r_k = p.add_run(k + " ")
            r_k.bold = True
            r_k.font.size = Pt(9.5)
            r_k.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            r_v = p.add_run(v)
            r_v.font.size = Pt(9.5)
            r_v.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 1
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Executive Scope & Design Principles Alignment")
    h1_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    p = doc.add_paragraph(
        "This document details the engineering implementation plan for the Major Project Review 2 Prototype of AEGIS. "
        "Based directly on the Functional Requirements Document (FRD) and system proposal, this prototype fulfills a strict design objective: "
        "deliver a high-fidelity, interactive, and visually stunning Frontend supported by a minimal, robust Backend executing REAL graph algorithms "
        "and REAL AI reasoning. Rather than static mock text, the prototype computes real blast radii, evaluates real composite risk formulas, executes real "
        "graph-grounded LLM prompts, and streams real-time reasoning tokens to the user interface."
    )
    p.paragraph_format.space_after = Pt(8)

    p_princ = doc.add_paragraph()
    p_princ.add_run("Core Design Principles Enforced in Prototype 1.0:\n").bold = True
    principles = [
        ("Unified Knowledge Substrate (FR-KG-1, FR-KG-2): ", "A single structural and runtime graph represents microservices, databases, APIs, and incident history."),
        ("Real Algorithmic Computation (FR-PS-1): ", "Blast radius and pre-ship risk scores are calculated via graph traversal (BFS / reachability algorithms) and weighted risk equations, not hardcoded dummy data."),
        ("Graph-Grounded LLM Reasoning (FR-MH-2): ", "The AI agent does not hallucinate over unstructured logs; it queries the graph topology neighborhood via Cypher/subgraph extraction to pinpoint root cause."),
        ("Strict Safety Invariant (FR-MH-4): ", "The AI reasoning layer is decoupled from production actuators. Candidate remediations are scored for safety, and execution is guarded by an automated Dead-Man's Rollback switch."),
        ("Closed-Loop Feedback Learning (FR-MH-6): ", "Healing outcomes are permanently committed back to the Knowledge Graph, dynamically updating future pre-ship risk weights.")
    ]
    for bold_txt, norm_txt in principles:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(bold_txt)
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        bp.add_run(norm_txt)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 2
    h2 = doc.add_heading(level=1)
    h2.add_run("2. System Architecture: The 'Frontend-First, Smart-Backend' Blueprint").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p = doc.add_paragraph(
        "To maximize visual polish and demonstration impact while keeping implementation overhead minimal, "
        "the architecture is partitioned into three lean tiers:"
    )
    
    arch_table = doc.add_table(rows=4, cols=3)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Tier", "Technology Choice", "Responsibility in Review 2 Prototype"]
    for c_idx, text in enumerate(headers):
        cell = arch_table.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tier_rows = [
        ("Presentation Layer (Frontend)", "React + TypeScript + TailwindCSS + @xyflow/react (React Flow) + Lucide Icons", "Renders the interactive simulation graph, dual-mode intelligence tabs, real-time edge traffic animations, blast radius halos, and streaming AI terminal."),
        ("Minimal Backend Layer (API & Algorithms)", "Python FastAPI + NetworkX / Neo4j Bolt Driver + Asyncio", "Serves graph topology, executes graph BFS blast radius traversals, computes pre-ship risk scores, coordinates chaos faults, and provides Server-Sent Events (SSE) streaming for AI tokens."),
        ("Reasoning & Memory Layer (AI & Graph)", "LiteLLM / OpenAI / Gemini API + Neo4j (Dockerized) / NetworkX Seed", "Generates graph Cypher queries, analyzes 2-hop fault neighborhoods, evaluates candidate fixes, and commits (:Incident) nodes back to the graph.")
    ]
    for r_idx, (t, tech, resp) in enumerate(tier_rows, start=1):
        for c_idx, val in enumerate([t, tech, resp]):
            cell = arch_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 3: REAL ALGORITHMS
    h3 = doc.add_heading(level=1)
    h3.add_run("3. Real Algorithms Specification").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph(
        "The prototype does NOT use hardcoded mockup calculations. The minimal backend implements three formal algorithms:"
    )

    doc.add_heading("Algorithm 1: Downstream Blast Radius & Impact Reachability (BFS Traversal)", level=2)
    p = doc.add_paragraph(
        "When a pull request modifies a target service or API contract, the system traverses all transitive downstream dependents "
        "across typed CALLS and DEPENDS_ON edges using Breadth-First Search with depth attenuation:"
    )
    code1 = doc.add_paragraph()
    code1.paragraph_format.left_indent = Inches(0.3)
    r = code1.add_run(
        "def compute_blast_radius(graph, changed_node_id, max_depth=3):\n"
        "    impacted = {changed_node_id: {'depth': 0, 'impact_weight': 1.0}}\n"
        "    queue = [(changed_node_id, 0, 1.0)]\n"
        "    while queue:\n"
        "        curr, depth, weight = queue.pop(0)\n"
        "        if depth < max_depth:\n"
        "            # Find services that CALL or DEPEND ON curr\n"
        "            for upstream in graph.predecessors(curr):\n"
        "                if upstream not in impacted:\n"
        "                    attenuated_weight = weight * 0.75\n"
        "                    impacted[upstream] = {'depth': depth + 1, 'impact_weight': attenuated_weight}\n"
        "                    queue.append((upstream, depth + 1, attenuated_weight))\n"
        "    return impacted"
    )
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_heading("Algorithm 2: Pre-Deployment Composite Risk Scoring Formula", level=2)
    p = doc.add_paragraph(
        "The overall pre-ship Risk Score R ∈ [0, 100] is computed mathematically as a weighted combination of structural blast radius, "
        "breaking contract severity, circular dependency penalty, and historical failure frequency:"
    )
    p_math = doc.add_paragraph()
    p_math.paragraph_format.left_indent = Inches(0.3)
    rm = p_math.add_run("RiskScore = w_b · B + w_c · C + w_d · D + w_h · H\n")
    rm.bold = True
    rm.font.name = 'Consolas'
    rm.font.size = Pt(10)
    rm.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

    p_terms = doc.add_paragraph(
        "Where:\n"
        "• B (Blast Radius Index): Ratio of impacted services to total system services (normalized to 0-100), weight w_b = 0.35.\n"
        "• C (Contract Breaking Severity): 100 if API endpoint payload/response schema is non-backward-compatible; 0 otherwise, weight w_c = 0.30.\n"
        "• D (Circular Dependency Penalty): 100 if the change introduces a directed cycle detected via Tarjan's strongly connected components; 0 otherwise, weight w_d = 0.15.\n"
        "• H (Historical Incident Factor): Number of incidents tied to the target node in the Knowledge Graph over the past 30 days (scaled 0-100), weight w_h = 0.20."
    )
    p_terms.paragraph_format.left_indent = Inches(0.2)

    doc.add_heading("Algorithm 3: Candidate Fix Simulation & Safety Scoring", level=2)
    p = doc.add_paragraph(
        "During autonomous healing, candidate remediations are scored before execution. For each fix candidate k:\n"
        "• Safety Index S_k = 100 - (BlastRisk(k) + DataLossRisk(k) + OverheadRisk(k))\n"
        "• Recovery Probability P_k = Similarity(CurrentAnomaly, HistoricalIncident_k)\n"
        "• Final Fix Score F_k = 0.6 · S_k + 0.4 · P_k. The actuator executes the candidate with max(F_k) provided S_k ≥ 75."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 4: REAL AI INTEGRATION
    h4 = doc.add_heading(level=1)
    h4.add_run("4. Real AI Integration & Graph-Grounded LLM Pipeline").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p = doc.add_paragraph(
        "Rather than simulating AI responses with pre-typed text, the minimal backend implements a live LLM orchestration pipeline "
        "connected to an LLM API (Gemini / OpenAI / Claude or local Ollama). The pipeline operates via three concrete phases:"
    )

    p_ai_phases = [
        ("Phase 1: Subgraph Neighborhood Extraction", "When an anomaly event fires on PaymentService, the backend queries the Knowledge Graph for the 2-hop subgraph: active connections, P99 latencies, error rates, database schema versions, and previous incident nodes. This structured context is formatted as a compact JSON graph payload."),
        ("Phase 2: Graph-Grounded RCA Prompting", "The LLM is prompted with a strict system prompt instructing it to reason exclusively over the provided graph topology. It is forbidden from guessing external causes not present in the graph. The agent produces: (1) Root Cause Hypothesis with confidence percentage, (2) Structural explanation of failure propagation, and (3) Ranked remediation recommendations."),
        ("Phase 3: Real-Time SSE Token Streaming to UI", "The backend streams the LLM's response to the frontend using Server-Sent Events (SSE). The frontend renders the agent's thought process live with a typewriter effect, displaying Cypher query lookups and reasoning steps in real-time.")
    ]
    for title, desc in p_ai_phases:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(title + ": ")
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        bp.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 5: MINIMAL BACKEND API
    h5 = doc.add_heading(level=1)
    h5.add_run("5. Minimal Backend Architecture (FastAPI Blueprint)").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p = doc.add_paragraph(
        "The backend is contained within a clean, single-purpose Python FastAPI service. It exposes 6 lightweight endpoints:"
    )

    api_table = doc.add_table(rows=7, cols=3)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Method & Endpoint", "Parameters", "Purpose & Algorithm Invoked"]
    for c_idx, text in enumerate(api_headers):
        cell = api_table.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    endpoints = [
        ("GET /api/graph/topology", "None", "Returns full 54-node, 142-edge ecosystem in React Flow JSON format."),
        ("POST /api/simulate-pr", "{pr_id, target_node, contract_diff}", "Executes Algorithm 1 (BFS Blast Radius) and Algorithm 2 (Composite Risk Score). Returns risk metrics."),
        ("POST /api/chaos/inject", "{fault_type, target_node}", "Updates runtime graph telemetry (e.g. latency -> 3500ms, status -> degraded). Triggers anomaly event."),
        ("GET /api/heal/stream", "{incident_id}", "SSE streaming endpoint. Runs Graph-Grounded LLM RCA Agent and streams reasoning tokens to UI."),
        ("POST /api/actuator/execute", "{candidate_id, guarded: true}", "Simulates fix execution, verifies health probe after 2s delay. Handles dead-man auto-rollback if probe fails."),
        ("POST /api/memory/commit", "{incident_id, rca, outcome}", "Creates new (:Incident) node in Knowledge Graph. Permanently updates future risk scoring weights.")
    ]
    for r_idx, (ep, params, desc) in enumerate(endpoints, start=1):
        for c_idx, val in enumerate([ep, params, desc]):
            cell = api_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.name = 'Consolas'
                r.font.color.rgb = RGBColor(0x37, 0x30, 0xA3)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 6: FRONTEND ARCHITECTURE
    h6 = doc.add_heading(level=1)
    h6.add_run("6. Frontend Architecture & Component Hierarchy").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p = doc.add_paragraph(
        "The frontend is constructed using React 18, Vite, and TailwindCSS. The component hierarchy is structured for maximum modularity and visual polish:"
    )

    fe_components = [
        ("DashboardHeader.tsx", "Displays the AEGIS logo, system status badge, the project motto banner, and the global KPI stats bar."),
        ("SimulationGraphCanvas.tsx", "Powered by React Flow (@xyflow/react). Features custom node designs (ServiceNode, DatabaseNode, PodNode), animated SVG edge pulses for traffic, and glowing red/amber halos for blast radius."),
        ("PreShipSimulatorPanel.tsx (Tab 1)", "Hosts the Pull Request selector, 'Run Pre-Ship Simulation' trigger, real-time risk gauge, contract drift diff viewer, and the CI/CD Merge Gate badge."),
        ("AutonomousHealingPanel.tsx (Tab 2)", "Houses the chaos injection dropdown, streaming AI reasoning terminal with typewriter effect, candidate fix cards, execution progress bar, and Dead-Man's switch toggle."),
        ("TimelineScrubber.tsx", "Interactive slider at the screen bottom allowing scrub-back replay across Baseline, Fault Injected, Healing, and Learned states."),
        ("IncidentMemoryBadge.tsx", "Animates the creation and attachment of new (:Incident) nodes onto the graph canvas upon incident resolution.")
    ]
    for comp, desc in fe_components:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(comp + " — ")
        r1.bold = True
        r1.font.name = 'Consolas'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        bp.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 7: FUTURE INTEGRATION ROADMAP
    h7 = doc.add_heading(level=1)
    h7.add_run("7. Future Integration Roadmap (Beyond Review 2)").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p = doc.add_paragraph(
        "While Review 2 proves the core algorithms, visual simulation, and AI reasoning on a controlled testbed, "
        "the final delivery for Review 3 will connect this foundation to full production cloud systems:"
    )

    roadmap_items = [
        ("Kubernetes Operator & In-Cluster Daemon (Phase 3)", "Replace simulated pod restarts with a native Go/Python Kubernetes operator listening on custom CRDs to perform rolling pod restarts and replica scaling directly on a live Kind/Minikube cluster."),
        ("Live Git Webhook & PR Bot Integration (Phase 3)", "Deploy a GitHub Actions / Webhook listener that intercepts real Pull Requests, runs the blast radius algorithm in CI/CD, and comments on the PR with the AEGIS Shield risk badge."),
        ("OpenTelemetry & eBPF Telemetry Ingest (Phase 4)", "Replace emulated metrics with real OpenTelemetry collectors and eBPF network sniffers to map service-to-service calls dynamically in production."),
        ("Chaos Mesh & Litmus Chaos Harness (Phase 4)", "Automate chaos injection directly inside Kubernetes to benchmark Time-to-Detect (TTD) and Time-to-Heal (TTH) across 50+ randomized microservice failure modes.")
    ]
    for r_title, r_desc in roadmap_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(r_title + ": ")
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        bp.add_run(r_desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 8: 5-DAY IMPLEMENTATION SCHEDULE
    h8 = doc.add_heading(level=1)
    h8.add_run("8. 5-Day Implementation Schedule & Milestones").font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    sched_table = doc.add_table(rows=6, cols=3)
    sched_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Day / Milestone", "Primary Tasks", "Verification Deliverable"]
    for c_idx, text in enumerate(s_headers):
        cell = sched_table.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    schedule = [
        ("Day 1: Topology & Canvas Setup", "Build React Flow canvas, create custom node designs (Service, DB, Pod), configure dark mode, render 54-node ecosystem.", "Interactive zoomable graph canvas rendering all nodes and styled connections."),
        ("Day 2: Minimal FastAPI & Real Algorithms", "Build FastAPI backend; implement Algorithm 1 (BFS Blast Radius) and Algorithm 2 (Composite Risk Score); connect to frontend.", "Clicking 'Run Simulation' on PR #142 computes real risk score and highlights affected downstream nodes."),
        ("Day 3: Real LLM Integration & Streaming", "Wire up LLM API with graph-grounded prompt template; build SSE streaming endpoint (/api/heal/stream); create frontend streaming terminal.", "Clicking 'Trigger Failure' streams real-time AI reasoning tokens onto the screen."),
        ("Day 4: Guarded Actuator & Closed Loop", "Implement candidate fix scoring, guarded execution delay, Dead-Man's switch fallback, and (:Incident) node memory injection.", "System heals in ~7s, renders auto-postmortem card, and updates Neo4j memory node."),
        ("Day 5: Replay Scrubber & Review 2 Polish", "Implement the timeline scrubber, polish UI animations (traffic dots, glowing halos), and rehearse the 4-minute demo walkthrough.", "Full Review 2 prototype ready for live demonstration to project committee.")
    ]
    for r_idx, (day, tasks, deliv) in enumerate(schedule, start=1):
        for c_idx, val in enumerate([day, tasks, deliv]):
            cell = sched_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x37, 0x30, 0xA3)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Final signoff
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sign = p_sign.add_run("AEGIS Engineering Project Team • Approved for Review 2 Prototype Development")
    r_sign.font.size = Pt(9)
    r_sign.font.italic = True
    r_sign.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Save
    doc.save("c:/Users/Fawaz/Desktop/AGEIS/AEGIS_Prototype_Implementation_Plan.docx")
    print("Implementation plan DOCX successfully created.")

if __name__ == "__main__":
    create_document()
